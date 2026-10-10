"""Uploaded standards documents, read and restated in one uniform shape.

Every organisation writes its contract positions differently: a Word memo
with numbered paragraphs, a spreadsheet, a scanned policy, a CSV in the Legal
template. Whatever arrives, each position comes out as the same record --
the fields of a StandardVersion, filed under one clause category -- so the
library, the analysis and the reports treat an uploaded standard exactly like
any other.

Two ways in:

    a table in the Legal template's columns     read directly, no model
    (CSV, XLSX, JSON)
    any other document (PDF, DOCX, scan, text)  read, split into parts, each
                                                part restated by the model

A model's reading is never trusted on its word. Each position carries the
sentence of the document it came from (`source_quote`); the quote is looked
for in the document's own text, and a position whose quote is not there is
left unticked for the reviewer. Nothing is saved until a person has reviewed
the positions, and what is saved arrives as drafts: a standard drives an
analysis only after someone with approval rights approves it (`approve`).
"""
from __future__ import annotations

import csv
import hashlib
import io
import json
import logging
import re
from datetime import timedelta
from pathlib import Path

from django.conf import settings
from django.core.exceptions import PermissionDenied, ValidationError
from django.db import transaction
from django.utils import timezone
from django.utils.translation import gettext as _

from apps.accounts import permissions as perms
from apps.accounts.models import Country
from apps.audit import services as audit
from apps.audit.models import AuditAction

from .models import (
    ClauseCategory,
    Standard,
    StandardProfile,
    StandardScope,
    StandardStatus,
    StandardUpload,
    StandardVersion,
)

logger = logging.getLogger(__name__)

MAX_MB = 20
#: Characters of a document given to the model at a time, and at most how many
#: parts: a standards document is tens of pages, not hundreds.
PART_CHARS = 6000
MAX_PARTS = 15

#: The uniform record every position is restated as.
TEXT_FIELDS = (
    "preferred_position", "preferred_position_ar", "acceptable_range", "red_flags",
    "approved_fallback_wording", "exceptions", "approval_authority",
    "responsible_department",
)
ITEM_FIELDS = ("category", *TEXT_FIELDS, "impact", "rule", "source_quote", "source_section",
               "country")

#: Spreadsheet / CSV / JSON column -> field. The Legal template's names, and
#: the fields' own names.
COLUMNS = {
    "category": "category", "clause_category": "category", "البند": "category", "الفئة": "category",
    "country": "country", "الدولة": "country",
    "preferred_position": "preferred_position", "position": "preferred_position",
    "الموقف_المفضل": "preferred_position",
    "preferred_position_ar": "preferred_position_ar",
    "acceptable_fallback": "acceptable_range", "acceptable_range": "acceptable_range",
    "acceptable": "acceptable_range", "المقبول": "acceptable_range",
    "walk_away": "red_flags", "red_flags": "red_flags", "red_line": "red_flags",
    "خط_أحمر": "red_flags",
    "approved_fallback_wording": "approved_fallback_wording", "fallback_wording": "approved_fallback_wording",
    "exceptions": "exceptions",
    "approval_authority": "approval_authority", "who_approves_the_fallback": "approval_authority",
    "responsible_department": "responsible_department", "department": "responsible_department",
    "impact": "impact", "impact_1_5": "impact",
    "legal_basis": "source_section", "source_reference": "source_section", "section": "source_section",
}

_RULE_EXTRACTS = {"percentage", "money", "duration_days"}
_RULE_OPERATORS = {"lte", "gte", "lt", "gt", "eq"}


# ── reading the file ─────────────────────────────────────────────────────

def _kind(data: bytes, filename: str) -> str:
    name = (filename or "").lower()
    if data.startswith(b"%PDF"):
        return "pdf"
    if data.startswith(b"PK"):
        if name.endswith(".xlsx"):
            return "xlsx"
        if name.endswith(".docx"):
            return "docx"
        raise ValidationError(_("نوع الملف غير مدعوم. المسموح: PDF أو Word أو Excel أو CSV أو JSON أو نص."))
    for ext in ("csv", "json", "txt", "md"):
        if name.endswith("." + ext):
            return ext
    raise ValidationError(_("نوع الملف غير مدعوم. المسموح: PDF أو Word أو Excel أو CSV أو JSON أو نص."))


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1256", "latin-1"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", "replace")


def _pdf_text(path: Path) -> tuple[str, list[str]]:
    """Native text where the PDF has it; Tesseract for scanned pages."""
    from apps.documents import extractors

    notes: list[str] = []
    parsed = extractors.extract(path, content_type="application/pdf")
    engine = None
    pages = []
    for page in parsed.pages:
        if page.has_native_text:
            pages.append(page.text)
            continue
        try:
            if engine is None:
                from apps.documents.ocr import get_ocr_engine

                engine = get_ocr_engine()
                ok, detail = engine.is_available()
                if not ok:
                    raise RuntimeError(detail)
            image = extractors.rasterise_page(path, page.page_number)
            result = engine.recognise_page(image, page_number=page.page_number,
                                           languages=settings.OCR_LANGUAGES)
            pages.append(result.text)
        except Exception as exc:
            notes.append(f"page {page.page_number}: not read ({exc})")
    return "\n\n".join(p for p in pages if p.strip()), notes


def _header_key(value) -> str:
    return re.sub(r"[\s\-/()]+", "_", str(value or "").strip().lower()).strip("_")


def _table_rows(rows: list[dict]) -> list[dict] | None:
    """Rows in the Legal template's columns, as items -- or None when the
    table is not in that shape and has to be read like a document."""
    if not rows:
        return None
    mapped_rows = []
    for row in rows:
        mapped: dict = {}
        for key, value in row.items():
            field = COLUMNS.get(_header_key(key))
            if field and value not in (None, ""):
                mapped[field] = str(value).strip() if not isinstance(value, (int, float)) else value
        if mapped:
            mapped_rows.append(mapped)
    if not mapped_rows or not all("category" in r and "preferred_position" in r
                                  for r in mapped_rows[:3]):
        return None
    return mapped_rows


def _xlsx_rows(path: Path) -> list[dict]:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    rows: list[dict] = []
    try:
        for sheet in book.worksheets:
            values = list(sheet.iter_rows(values_only=True))
            header_at = next((i for i, r in enumerate(values[:10])
                              if r and any(COLUMNS.get(_header_key(c)) for c in r if c)), None)
            if header_at is None:
                continue
            header = [str(c or "") for c in values[header_at]]
            for r in values[header_at + 1:]:
                if r and any(c not in (None, "") for c in r):
                    rows.append(dict(zip(header, r)))
    finally:
        book.close()
    return rows


def _xlsx_text(path: Path) -> str:
    from openpyxl import load_workbook

    book = load_workbook(path, read_only=True, data_only=True)
    lines = []
    try:
        for sheet in book.worksheets:
            lines.append(f"[{sheet.title}]")
            for r in sheet.iter_rows(values_only=True):
                cells = [str(c).strip() for c in r if c not in (None, "")]
                if cells:
                    lines.append(" | ".join(cells))
    finally:
        book.close()
    return "\n".join(lines)


def read(path: Path, filename: str) -> dict:
    """{"table": [rows] | None, "text": str, "notes": [...]}."""
    data = path.read_bytes()
    kind = _kind(data, filename)
    notes: list[str] = []
    table, text = None, ""
    if kind == "pdf":
        text, notes = _pdf_text(path)
    elif kind == "docx":
        from apps.documents import extractors

        parsed = extractors.extract(
            path, content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document")
        text = "\n".join(p.text for p in parsed.pages)
    elif kind == "xlsx":
        table = _table_rows(_xlsx_rows(path))
        text = _xlsx_text(path)
    elif kind == "csv":
        raw = _decode(data)
        table = _table_rows(list(csv.DictReader(io.StringIO(raw))))
        text = raw
    elif kind == "json":
        raw = _decode(data)
        try:
            loaded = json.loads(raw)
        except json.JSONDecodeError as exc:
            raise ValidationError(_("ملف JSON غير صالح.")) from exc
        if isinstance(loaded, dict):
            loaded = loaded.get("standards") or loaded.get("positions") or [loaded]
        table = _table_rows([r for r in loaded if isinstance(r, dict)]) if isinstance(loaded, list) else None
        text = raw
    else:
        text = _decode(data)
    return {"table": table, "text": text.strip(), "notes": notes}


# ── the uniform record ───────────────────────────────────────────────────

def categories() -> list[dict]:
    return [
        {"code": c.code, "name": c.name, "name_ar": c.name_ar, "description": c.description}
        for c in ClauseCategory.objects.filter(is_active=True)
    ]


def _category_code(value: str, known: list[dict]) -> str:
    """A category given as its code, its English name or its Arabic name."""
    wanted = (value or "").strip().casefold()
    if not wanted:
        return ""
    for c in known:
        if wanted in {c["code"].casefold(), c["name"].casefold(), (c["name_ar"] or "").casefold()}:
            return c["code"]
    for c in known:                           # "Liability cap" for "Liability cap and indemnities"
        if len(wanted) >= 5 and (wanted in c["name"].casefold() or wanted in (c["name_ar"] or "").casefold()):
            return c["code"]
    return ""


def _country_code(value) -> str:
    wanted = str(value or "").strip().casefold()
    if not wanted or wanted in {"global", "all", "*", "عام", "الكل"}:
        return ""
    for code, label in Country.choices:
        if wanted in {code.casefold(), str(label).casefold()}:
            return code
    aliases = {"ksa": "SA", "saudi": "SA", "السعودية": "SA", "الأردن": "JO", "jordan": "JO",
               "uae": "AE", "الإمارات": "AE", "عمان": "OM", "عُمان": "OM"}
    return aliases.get(wanted, "")


def clean_rule(rule) -> dict:
    """A proposed rule the rules engine can evaluate, or {}.

    The engine evaluates rules in Python (apps.analysis.rules); a shape it
    does not understand would silently check nothing, so it is dropped here
    and the reviewer can add one by hand.
    """
    if not isinstance(rule, dict) or not rule:
        return {}
    operator = rule.get("operator")
    if operator in {"present", "absent"}:
        markers = [str(m).strip() for m in rule.get("markers") or [] if str(m).strip()]
        return {"operator": operator, "markers": markers[:12]} if markers else {}
    if operator not in _RULE_OPERATORS or rule.get("extract") not in _RULE_EXTRACTS:
        return {}
    try:
        out = {"extract": rule["extract"], "operator": operator, "value": float(rule["value"])}
        if rule.get("fallback_value") not in (None, ""):
            out["fallback_value"] = float(rule["fallback_value"])
    except (KeyError, TypeError, ValueError):
        return {}
    for key in ("value", "fallback_value"):
        if key in out and out[key].is_integer():
            out[key] = int(out[key])
    if rule["extract"] == "money" and rule.get("currency"):
        out["currency"] = str(rule["currency"]).upper()[:3]
    return out


def quote_found(quote: str, text: str) -> bool | None:
    """Is the quoted sentence in the document? None when there is no quote.

    Compared after the normalisation OCR verification uses (spacing,
    diacritics, alef forms, digit scripts), so an honest quote of a scanned
    page still matches. A quote made of pieces joined by "…" must have every
    piece in the document.
    """
    from apps.documents.verification import normalise_for_compare

    quote = (quote or "").strip()
    if not quote:
        return None
    haystack = normalise_for_compare(text)
    pieces = [normalise_for_compare(p) for p in re.split(r"\s*(?:…|\.\.\.)\s*", quote)]
    pieces = [p for p in pieces if len(p) >= 8]
    return bool(pieces) and all(p in haystack for p in pieces)


def _item(raw: dict, known: list[dict], text: str, default_country: str) -> dict:
    item = {field: "" for field in ITEM_FIELDS}
    for field in TEXT_FIELDS + ("source_quote", "source_section"):
        item[field] = str(raw.get(field) or "").strip()
    item["category"] = _category_code(str(raw.get("category") or ""), known)
    item["country"] = _country_code(raw.get("country")) if raw.get("country") else default_country
    try:
        impact = int(raw.get("impact")) if raw.get("impact") not in (None, "") else None
    except (TypeError, ValueError):
        impact = None
    item["impact"] = impact if impact in {1, 2, 3, 4, 5} else None
    item["rule"] = clean_rule(raw.get("rule") or {})
    item["quote_found"] = quote_found(item["source_quote"], text)
    problems = []
    if not item["category"]:
        problems.append("unknown_category")
    if not item["preferred_position"]:
        problems.append("no_position")
    if item["quote_found"] is False:
        problems.append("quote_not_found")
    if raw.get("rule") and not item["rule"]:
        problems.append("rule_dropped")
    item["problems"] = problems
    # Ticked only when nothing is wrong; a person can still tick the rest.
    item["include"] = not problems
    return item


def _parts(text: str) -> list[str]:
    """Paragraph-aligned parts of at most PART_CHARS."""
    parts, current = [], ""
    for para in re.split(r"\n\s*\n", text):
        para = para.strip()
        if not para:
            continue
        while len(para) > PART_CHARS:                 # one enormous paragraph
            if current:
                parts.append(current); current = ""
            parts.append(para[:PART_CHARS]); para = para[PART_CHARS:]
        if current and len(current) + len(para) + 2 > PART_CHARS:
            parts.append(current); current = ""
        current = f"{current}\n\n{para}" if current else para
    if current:
        parts.append(current)
    return parts


def _merge(items: list[dict]) -> list[dict]:
    """One entry per category and country: the model sees the document in
    parts, so one category can be found in two of them."""
    merged: dict[tuple, dict] = {}
    order = []
    for item in items:
        key = (item["category"], item["country"]) if item["category"] else (id(item),)
        if key not in merged:
            merged[key] = item
            order.append(key)
            continue
        base = merged[key]
        for field in TEXT_FIELDS:
            if item[field] and item[field] not in base[field]:
                base[field] = f"{base[field]}\n{item[field]}".strip()
        if item["source_quote"] and item["source_quote"] not in base["source_quote"]:
            base["source_quote"] = f"{base['source_quote']} … {item['source_quote']}".strip(" …")
        base["source_section"] = base["source_section"] or item["source_section"]
        base["rule"] = base["rule"] or item["rule"]
        base["impact"] = base["impact"] or item["impact"]
        base["quote_found"] = (None if base["quote_found"] is None and item["quote_found"] is None
                               else bool(base["quote_found"] is not False and item["quote_found"] is not False))
        base["problems"] = sorted(set(base["problems"]) | set(item["problems"]))
        base["include"] = not base["problems"]
    return [merged[k] for k in order]


# ── the pipeline ─────────────────────────────────────────────────────────

def create_upload(user, data: bytes, filename: str, *, profile: str, country: str,
                  effective_date=None, note: str = "") -> StandardUpload:
    """Store an upload and read it. Raises ValidationError for a bad file."""
    from apps.documents.storage import safe_filename

    perms.require_capability(user, perms.MANAGE_STANDARDS)
    if len(data) > MAX_MB * 1024 * 1024:
        raise ValidationError(_("الملف أكبر من الحد المسموح (%(mb)s ميغابايت).") % {"mb": MAX_MB})
    _kind(data, filename)                                     # refuses what cannot be read
    if profile not in StandardProfile.values:
        raise ValidationError(_("مرجع المعايير غير معروف."))
    if profile == StandardProfile.FIDIC:
        # FIDIC is an international form: one position for every country.
        # Country-specific positions are AJi's own (they take precedence over
        # FIDIC country by country), so a country posted with FIDIC is dropped.
        country = ""
    if country and country not in dict(Country.choices):
        raise ValidationError(_("الدولة غير معروفة."))
    digest = hashlib.sha256(data).hexdigest()
    name = safe_filename(filename or "standards")[:200]
    relative = f"standards/{digest}{Path(name).suffix.lower()[:8]}"
    target = Path(settings.PRIVATE_MEDIA_ROOT) / relative
    target.parent.mkdir(parents=True, exist_ok=True)
    if not target.exists():
        target.write_bytes(data)
    upload = StandardUpload.objects.create(
        uploaded_by=user, original_filename=name, stored_path=relative, sha256=digest,
        size_bytes=len(data), profile=profile, country=country or "",
        effective_date=effective_date or timezone.localdate(), note=(note or "")[:500],
    )
    process(upload)
    upload.refresh_from_db()
    return upload


def process(upload: StandardUpload) -> StandardUpload:
    """Read the file and restate its positions. Never raises: a failure is
    recorded on the upload."""
    from apps.ai.providers.base import AIProviderError
    from apps.ai.services import ai_service

    known = categories()
    upload.status = StandardUpload.Status.PROCESSING
    upload.error = ""
    upload.ai_calls = 0
    try:
        path = Path(settings.PRIVATE_MEDIA_ROOT) / upload.stored_path
        found = read(path, upload.original_filename)
        text = found["text"]
        upload.extracted_text = text
        notes = list(found["notes"])
        if found["table"] is not None:
            upload.method = "table"
            # A table's cells are the document; each row quotes itself.
            items = []
            if upload.profile == StandardProfile.FIDIC:
                for row in found["table"]:
                    row.pop("country", None)
            for row in found["table"]:
                row = dict(row)
                row.setdefault("source_quote", "")
                items.append(_item(row, known, text, upload.country))
        else:
            if not text:
                raise ValueError(_("لم يُعثر على نص مقروء في الملف."))
            upload.method = "ai"
            parts = _parts(text)
            if len(parts) > MAX_PARTS:
                notes.append(f"only the first {MAX_PARTS} of {len(parts)} parts were read")
                parts = parts[:MAX_PARTS]
            items, failures = [], 0
            for n, part in enumerate(parts, 1):
                try:
                    call = ai_service.normalize_standard(text=part, categories=known, part=n, parts=len(parts))
                except AIProviderError as exc:
                    failures += 1
                    notes.append(f"part {n}: {type(exc).__name__}")
                    continue
                finally:
                    upload.ai_calls += 1
                items += [_item(p.model_dump(), known, text, upload.country) for p in call.parsed.positions]
                if call.parsed.notes:
                    notes.append(f"part {n}: {call.parsed.notes[:300]}")
            if failures == len(parts):
                raise RuntimeError(_("تعذّر الوصول إلى خدمة الذكاء الاصطناعي لقراءة الملف."))
        upload.items = _merge(items)
        upload.status = StandardUpload.Status.REVIEW
        upload.error = "\n".join(notes)[:4000]
    except (ValidationError, ValueError, RuntimeError, OSError) as exc:
        message = exc.messages[0] if isinstance(exc, ValidationError) else str(exc)
        upload.status = StandardUpload.Status.FAILED
        upload.error = message[:4000]
    except Exception as exc:                 # a broken file must not become a 500
        logger.exception("Standards upload %s could not be read", upload.pk)
        upload.status = StandardUpload.Status.FAILED
        upload.error = f"{type(exc).__name__}: {exc}"[:4000]
    upload.processed_at = timezone.now()
    upload.save()
    return upload


def items_from_post(post, count: int) -> list[dict]:
    """The reviewer's edited positions, from the review form."""
    known = categories()
    out = []
    for i in range(count):
        p = f"i{i}-"
        raw = {field: post.get(p + field, "") for field in TEXT_FIELDS + ("source_quote", "source_section")}
        raw["category"] = post.get(p + "category", "")
        raw["country"] = post.get(p + "country", "")
        raw["impact"] = post.get(p + "impact", "")
        rule_text = (post.get(p + "rule") or "").strip()
        try:
            raw["rule"] = json.loads(rule_text) if rule_text else {}
        except json.JSONDecodeError:
            raw["rule"] = {"invalid": rule_text}
        item = _item(raw, known, "", "")
        item["country"] = _country_code(raw["country"]) if raw["country"] else ""
        item["quote_found"] = None
        item["include"] = post.get(p + "include") == "1"
        out.append(item)
    return out


@transaction.atomic
def save_drafts(upload: StandardUpload, items: list[dict], user) -> list[StandardVersion]:
    """The ticked positions, as new versions awaiting approval.

    Each goes onto the standard for its category, scope, country and profile
    -- a new version beside whatever is approved there now, which keeps
    applying until this one is approved.
    """
    perms.require_capability(user, perms.MANAGE_STANDARDS)
    by_code = {c.code: c for c in ClauseCategory.objects.all()}
    created = []
    for item in items:
        if not item.get("include"):
            continue
        category = by_code.get(item["category"])
        if category is None or not item["preferred_position"]:
            continue
        country = item.get("country") or upload.country
        if upload.profile == StandardProfile.FIDIC:
            country = ""                  # a spreadsheet row cannot make FIDIC national
        scope = StandardScope.COUNTRY if country else StandardScope.GLOBAL
        standard, _new = Standard.objects.get_or_create(
            category=category, scope=scope, country=country, contract_type=None,
            reference_template=upload.profile,
            defaults={"title": f"{category.name} — {country or 'global'} ({upload.get_profile_display()})"},
        )
        if not standard.is_active:
            standard.is_active = True
            standard.save(update_fields=["is_active", "updated_at"])
        latest = standard.versions.order_by("-version").first()
        section = f", {item['source_section']}" if item.get("source_section") else ""
        version = StandardVersion.objects.create(
            standard=standard, version=(latest.version + 1) if latest else 1,
            preferred_position=item["preferred_position"],
            preferred_position_ar=item.get("preferred_position_ar", ""),
            acceptable_range=item.get("acceptable_range", ""),
            red_flags=item.get("red_flags", ""),
            approved_fallback_wording=item.get("approved_fallback_wording", ""),
            exceptions=item.get("exceptions", ""),
            approval_authority=item.get("approval_authority", "")[:160],
            responsible_department=item.get("responsible_department", "")[:64],
            rule_definition=clean_rule(item.get("rule") or {}),
            impact=item.get("impact"),
            source_reference=f"Uploaded: {upload.original_filename}{section}"[:255],
            status=StandardStatus.PENDING_APPROVAL,
            effective_date=upload.effective_date,
            change_note=(f"{upload.note}\n" if upload.note else "")
            + (f"Source: {item['source_quote']}" if item.get("source_quote") else ""),
            is_demo_fixture=False,
        )
        audit.record(
            AuditAction.STANDARD_CREATED, actor=user, entity_type="StandardVersion",
            entity_id=version.pk,
            details={"standard": str(standard), "version": version.version,
                     "upload": str(upload.pk), "file": upload.original_filename},
        )
        created.append(version)
    upload.items = items
    upload.created_versions = [str(v.pk) for v in created]
    upload.status = StandardUpload.Status.SAVED
    upload.save(update_fields=["items", "created_versions", "status", "updated_at"])
    return created


@transaction.atomic
def approve(version: StandardVersion, user) -> StandardVersion:
    """Approve a draft: from its effective date it is the standard that
    applies, and the version it replaces is superseded (kept, so analyses
    that cited it stay explainable)."""
    perms.require_capability(user, perms.APPROVE_STANDARDS)
    if version.status not in (StandardStatus.DRAFT, StandardStatus.PENDING_APPROVAL):
        raise ValidationError(_("هذا الإصدار ليس مسودة بانتظار الاعتماد."))
    if not version.effective_date:
        raise ValidationError(_("لا يمكن اعتماد إصدار بلا تاريخ سريان."))
    if version.impact is None:
        raise ValidationError(_("حدّد درجة الأثر (1-5) قبل الاعتماد؛ هي ما يغذّي تقييم المخاطر."))
    for previous in version.standard.versions.filter(status=StandardStatus.APPROVED):
        previous.status = StandardStatus.SUPERSEDED
        previous.expiry_date = max(previous.effective_date or version.effective_date,
                                   version.effective_date - timedelta(days=1))
        previous.save(update_fields=["status", "expiry_date", "updated_at"])
        audit.record(
            AuditAction.STANDARD_SUPERSEDED, actor=user, entity_type="StandardVersion",
            entity_id=previous.pk, previous_state=StandardStatus.APPROVED,
            new_state=StandardStatus.SUPERSEDED,
            details={"standard": str(previous.standard), "version": previous.version,
                     "by_version": version.version},
        )
    version.status = StandardStatus.APPROVED
    version.approved_by = user
    version.approved_at = timezone.now()
    version.save(update_fields=["status", "approved_by", "approved_at", "updated_at"])
    audit.record(
        AuditAction.STANDARD_APPROVED, actor=user, entity_type="StandardVersion",
        entity_id=version.pk, new_state=StandardStatus.APPROVED,
        details={"standard": str(version.standard), "version": version.version},
    )
    return version


def set_impact(version: StandardVersion, impact: int, user) -> None:
    """Impact is Legal's to set, on a draft, before approval."""
    perms.require_capability(user, perms.MANAGE_STANDARDS)
    if version.is_locked:
        raise ValidationError(_("الإصدار المعتمد لا يُعدّل."))
    if impact not in {1, 2, 3, 4, 5}:
        raise ValidationError(_("درجة الأثر من 1 إلى 5."))
    version.impact = impact
    version.save(update_fields=["impact", "updated_at"])


def withdraw(version: StandardVersion, user) -> None:
    """Take a draft off the approval queue. Approved versions are superseded
    by a newer one, never withdrawn here."""
    perms.require_capability(user, perms.MANAGE_STANDARDS)
    if version.status not in (StandardStatus.DRAFT, StandardStatus.PENDING_APPROVAL):
        raise PermissionDenied(_("لا يمكن سحب إلا المسودة."))
    version.status = StandardStatus.WITHDRAWN
    version.save(update_fields=["status", "updated_at"])
