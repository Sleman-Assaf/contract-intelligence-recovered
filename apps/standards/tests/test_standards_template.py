"""The Excel template Legal fills in, and its import.

Round trip: write the template, fill it the way Legal would -- an English row,
an Arabic-only row, a Jordan override for one contract type, a real Excel
date -- and load it with import_standards. Untouched rows are skipped, the
Arabic columns survive, everything arrives as a draft."""
from __future__ import annotations

from datetime import datetime

import pytest
from django.core.management import call_command
from django.core.management.base import CommandError
from openpyxl import load_workbook

from apps.standards.models import ClauseCategory, StandardStatus, StandardVersion

pytestmark = pytest.mark.django_db


@pytest.fixture
def template(tmp_path, settings):
    settings.DEBUG = True
    call_command("seed_reference", verbosity=0)
    path = tmp_path / "AJi_Standards_Template.xlsx"
    call_command("standards_template", str(path), verbosity=0)
    return path


def _columns(ws):
    return {ws.cell(row=3, column=c).value: c for c in range(1, ws.max_column + 1)}


def _row_for(ws, code):
    return next(r for r in range(4, ws.max_row + 1) if ws.cell(row=r, column=1).value == code)


class TestTheTemplate:
    def test_one_row_per_active_category(self, template):
        ws = load_workbook(template)["Standards"]
        codes = [ws.cell(row=r, column=1).value for r in range(4, ws.max_row + 1)]
        expected = set(ClauseCategory.objects.filter(is_active=True).values_list("code", flat=True))
        assert expected <= set(codes)

    def test_the_key_row_is_what_the_importer_reads(self, template):
        from apps.standards.management.commands.import_standards import CSV_FIELDS

        keys = _columns(load_workbook(template)["Standards"])
        for key in ("category", "country", "contract_type", "preferred_position", "preferred_position_ar",
                    "approved_fallback_wording", "approved_fallback_wording_ar", "impact", "effective_date"):
            assert key in keys and key in CSV_FIELDS

    def test_a_blank_template_imports_nothing(self, template):
        with pytest.raises(CommandError, match="No rows"):
            call_command("import_standards", str(template), verbosity=0)


class TestAFilledTemplateImports:
    def _fill(self, template):
        book = load_workbook(template)
        ws = book["Standards"]
        col = _columns(ws)
        r = _row_for(ws, "liability")
        ws.cell(row=r, column=col["preferred_position"], value="Aggregate liability capped at the fee.")
        ws.cell(row=r, column=col["approved_fallback_wording"], value="The Consultant's total liability shall not exceed the fee.")
        ws.cell(row=r, column=col["approved_fallback_wording_ar"], value="لا تتجاوز المسؤولية الإجمالية للاستشاري قيمة الأتعاب.")
        ws.cell(row=r, column=col["impact"], value=5)
        ws.cell(row=r, column=col["effective_date"], value=datetime(2027, 1, 1))
        r = _row_for(ws, "delayLiq")
        ws.cell(row=r, column=col["preferred_position_ar"], value="سقف تعويضات التأخير 10% من قيمة العقد.")
        ws.cell(row=r, column=col["impact"], value=4)
        extra = next(x for x in range(4, ws.max_row + 1) if not ws.cell(row=x, column=1).value)
        ws.cell(row=extra, column=1, value="dispute")
        ws.cell(row=extra, column=col["country"], value="JO")
        ws.cell(row=extra, column=col["contract_type"], value="supervision")
        ws.cell(row=extra, column=col["preferred_position"], value="Arbitration in Amman.")
        ws.cell(row=extra, column=col["impact"], value=3)
        book.save(template)

    def test_filled_rows_arrive_as_drafts_with_the_arabic_kept(self, template):
        self._fill(template)
        call_command("import_standards", str(template), verbosity=0)
        imported = StandardVersion.objects.filter(is_demo_fixture=False, standard__reference_template="aji_standard")
        assert imported.count() == 3
        assert set(imported.values_list("status", flat=True)) == {StandardStatus.DRAFT}

        liability = imported.get(standard__category__code="liability")
        assert liability.impact == 5
        assert liability.approved_fallback_wording_ar.startswith("لا تتجاوز")
        assert str(liability.effective_date) == "2027-01-01"

        delay = imported.get(standard__category__code="delayLiq")
        assert delay.preferred_position_ar.startswith("سقف")
        assert delay.preferred_position == delay.preferred_position_ar     # Arabic-only is enough

        jordan = imported.get(standard__category__code="dispute")
        assert jordan.standard.country == "JO"
        assert jordan.standard.contract_type.code == "supervision"

    def test_the_aji_role_column_is_kept(self, template):
        book = load_workbook(template)
        ws = book["Standards"]
        col = _columns(ws)
        r = _row_for(ws, "liability")
        ws.cell(row=r, column=col["preferred_position"], value="Sub-consultant's cap at least its fee.")
        ws.cell(row=r, column=col["impact"], value=5)
        ws.cell(row=r, column=col["aji_role"], value="client")
        book.save(template)
        call_command("import_standards", str(template), verbosity=0)
        imported = StandardVersion.objects.get(is_demo_fixture=False, standard__category__code="liability")
        assert imported.standard.perspective == "client"

    def test_fallback_and_red_lines_keep_both_languages(self, template):
        book = load_workbook(template)
        ws = book["Standards"]
        col = _columns(ws)
        r = _row_for(ws, "liability")
        ws.cell(row=r, column=col["preferred_position"], value="Cap at the fee.")
        ws.cell(row=r, column=col["acceptable_range"], value="Cap at the insurance cover.")
        ws.cell(row=r, column=col["acceptable_range_ar"], value="سقف يساوي التغطية التأمينية.")
        ws.cell(row=r, column=col["red_flags"], value="No cap.")
        ws.cell(row=r, column=col["red_flags_ar"], value="لا سقف للمسؤولية.")
        ws.cell(row=r, column=col["impact"], value=5)
        book.save(template)
        call_command("import_standards", str(template), verbosity=0)
        v = StandardVersion.objects.get(is_demo_fixture=False, standard__category__code="liability")
        assert (v.acceptable_range, v.acceptable_range_ar) == ("Cap at the insurance cover.", "سقف يساوي التغطية التأمينية.")
        assert (v.red_flags, v.red_flags_ar) == ("No cap.", "لا سقف للمسؤولية.")

    def test_a_bad_impact_fails_the_whole_file(self, template):
        self._fill(template)
        book = load_workbook(template)
        ws = book["Standards"]
        ws.cell(row=_row_for(ws, "liability"), column=_columns(ws)["impact"], value=9)
        book.save(template)
        with pytest.raises(CommandError, match="failed validation"):
            call_command("import_standards", str(template), verbosity=0)
        assert not StandardVersion.objects.filter(is_demo_fixture=False, standard__reference_template="aji_standard").exists()
