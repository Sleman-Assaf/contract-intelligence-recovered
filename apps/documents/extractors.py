"""Native extraction — the preferred path (§5, §6).

Order of preference is native text → OCR → vision, and it exists for two
reasons. Native text is exact, and it carries real coordinates, which is what
makes "open the source and highlight clause 12.3 on page 14" possible at all.

Text extracted here is preserved verbatim. No model rewrites it (§6, §53).
"""
from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from django.conf import settings

from .ocr.base import BBox
from .text import plain_letters


@dataclass
class ExtractedBlock:
    text: str
    bbox: BBox | None
    reading_order: int


@dataclass
class ExtractedPage:
    page_number: int
    blocks: list[ExtractedBlock] = field(default_factory=list)
    width: float | None = None
    height: float | None = None
    rotation: int = 0
    has_native_text: bool = False

    @property
    def text(self) -> str:
        return "\n".join(b.text for b in self.blocks)

    @property
    def char_count(self) -> int:
        return sum(len(b.text.strip()) for b in self.blocks)


@dataclass
class ExtractedDocument:
    pages: list[ExtractedPage] = field(default_factory=list)
    is_encrypted: bool = False
    page_count: int = 0

    @property
    def native_text_coverage(self) -> float:
        """Share of pages with a usable text layer."""
        if not self.pages:
            return 0.0
        return sum(1 for p in self.pages if p.has_native_text) / len(self.pages)

    @property
    def requires_ocr(self) -> bool:
        return self.native_text_coverage < settings.NATIVE_TEXT_MIN_COVERAGE


class ExtractionError(Exception):
    """Unreadable, corrupted or password-protected source (§43)."""


class EncryptedDocumentError(ExtractionError):
    """Password-protected. A human must supply an unlocked copy."""


def _min_chars_for_native_text() -> int:
    """Below this many characters a page is treated as image-only, whatever the
    PDF claims — scanned pages often carry a few stray header characters.

    Configuration, not a constant: the right value depends on the contracts an
    organisation actually receives (§50).
    """
    return getattr(settings, "NATIVE_TEXT_MIN_CHARS", 120)


def extract_pdf(path: Path | str) -> ExtractedDocument:
    """Native PDF text with per-block bounding boxes."""
    import pymupdf

    try:
        doc = pymupdf.open(str(path))
    except Exception as exc:
        raise ExtractionError(f"Could not open the PDF: {exc}") from exc

    try:
        if doc.needs_pass:
            raise EncryptedDocumentError(
                "This PDF is password-protected and cannot be analysed. "
                "Please upload an unlocked copy."
            )

        out = ExtractedDocument(page_count=doc.page_count)
        for index, page in enumerate(doc, start=1):
            rect = page.rect
            extracted = ExtractedPage(
                page_number=index,
                width=float(rect.width),
                height=float(rect.height),
                rotation=int(page.rotation or 0),
            )

            try:
                raw = page.get_text("dict")
            except Exception as exc:  # a damaged page must not kill the document
                extracted.blocks.append(
                    ExtractedBlock(text="", bbox=None, reading_order=0)
                )
                out.pages.append(extracted)
                continue

            order = 0
            for block in raw.get("blocks", []):
                if block.get("type") != 0:      # 0 = text, 1 = image
                    continue
                lines: list[str] = []
                for line in block.get("lines", []):
                    spans = [s.get("text", "") for s in line.get("spans", [])]
                    # Ligature glyphs as the letters they draw (documents/text.py).
                    joined = plain_letters("".join(spans)).strip()
                    if joined:
                        lines.append(joined)
                text = "\n".join(lines).strip()
                if not text:
                    continue
                bx = block.get("bbox")
                extracted.blocks.append(
                    ExtractedBlock(
                        text=text,
                        bbox=BBox(*(float(v) for v in bx)) if bx else None,
                        reading_order=order,
                    )
                )
                order += 1

            extracted.has_native_text = extracted.char_count >= _min_chars_for_native_text()
            out.pages.append(extracted)

        return out
    finally:
        doc.close()


#: The resolution pages are rendered at for OCR. Tesseract's boxes are pixels
#: of that image; everything stored is in PDF points (1/72 inch), which is what
#: `crop_region`, the source viewer and native text all use.
#: (Scanned contract pages now render at settings.OCR_RENDER_DPI through
#: apps/documents/ocr/pipeline.py; this stays the default for the chat and
#: standards readers.)
OCR_RENDER_DPI = 180
PDF_POINTS_PER_INCH = 72


def pixels_to_points(bbox, dpi: int = OCR_RENDER_DPI) -> tuple[float, float, float, float]:
    """An OCR box (pixels of a `dpi` render) in PDF points.

    `pdf_coord = ocr_coord * (72 / dpi)`. Without it every OCR box was stored
    2.5 times too large: measured on a scanned A4 page (595 x 842 pt), the
    "4.3 Delay Damages" line came back at y = 1832-1932 -- off the page -- so
    the vision crop failed outright, and boxes higher up cropped the wrong
    text.
    """
    scale = PDF_POINTS_PER_INCH / dpi
    x0, y0, x1, y1 = bbox
    return (x0 * scale, y0 * scale, x1 * scale, y1 * scale)


def _checked_box(bbox) -> tuple[float, float, float, float]:
    """Four finite numbers with a positive width and height, or ExtractionError."""
    import math

    try:
        x0, y0, x1, y1 = (float(v) for v in bbox)
    except (TypeError, ValueError) as exc:
        raise ExtractionError(f"Invalid region {bbox!r}.") from exc
    if not all(math.isfinite(v) for v in (x0, y0, x1, y1)):
        raise ExtractionError(f"Invalid region {bbox!r}.")
    if x1 <= x0 or y1 <= y0:
        raise ExtractionError(f"Empty region {bbox!r}.")
    return x0, y0, x1, y1


def rasterise_page(path: Path | str, page_number: int, *, dpi: int = OCR_RENDER_DPI) -> bytes:
    """Render one page to PNG for OCR or vision.

    180 DPI is about the floor at which Arabic diacritics survive; below it,
    OCR confidence on Arabic drops sharply.
    """
    import pymupdf

    doc = pymupdf.open(str(path))
    try:
        if page_number < 1 or page_number > doc.page_count:
            raise ExtractionError(f"Page {page_number} is out of range.")
        page = doc[page_number - 1]
        pix = page.get_pixmap(dpi=dpi)
        return pix.tobytes("png")
    finally:
        doc.close()


def crop_region(path: Path | str, page_number: int, bbox, *, dpi: int = 220,
                padding: float = 6.0, rotate: int = 0) -> bytes:
    """Render just one region — what vision verification actually needs (§10).

    Sending a cropped clause instead of the whole page is the difference
    between verifying a suspect number and re-reading a 200-page contract.

    `bbox` is in PDF points. A region that is invalid, empty, or lies
    outside the page raises ExtractionError -- never a renderer error, which
    used to escape the vision step uncaught and fail the whole extraction.
    """
    import pymupdf

    x0, y0, x1, y1 = _checked_box(bbox)
    try:
        doc = pymupdf.open(str(path))
    except Exception as exc:
        raise ExtractionError(f"Could not open the file to crop: {exc}") from exc
    try:
        if page_number < 1 or page_number > doc.page_count:
            raise ExtractionError(f"Page {page_number} is out of range.")
        page = doc[page_number - 1]
        rect = pymupdf.Rect(
            max(0, x0 - padding),
            max(0, y0 - padding),
            min(page.rect.width, x1 + padding),
            min(page.rect.height, y1 + padding),
        )
        # Less than a point either way is not a region with anything in it:
        # most often a box that was never on this page at all.
        if rect.is_empty or rect.width < 1 or rect.height < 1:
            raise ExtractionError(
                f"Region {bbox!r} lies outside page {page_number} "
                f"({page.rect.width:.0f} x {page.rect.height:.0f} pt).")
        try:
            pix = page.get_pixmap(dpi=dpi, clip=rect)
            png = pix.tobytes("png")
        except Exception as exc:
            raise ExtractionError(f"Could not render region {bbox!r}: {exc}") from exc
        if rotate % 360:
            # `rotate` is how far OCR turned the page clockwise to read it
            # (Page.ocr_rotation); the crop is turned the same way.
            import io

            from PIL import Image

            with Image.open(io.BytesIO(png)) as img:
                turned = img.rotate(-(rotate % 360), expand=True)
                out = io.BytesIO()
                turned.save(out, "PNG")
                png = out.getvalue()
        return png
    finally:
        doc.close()


def crop_image(image_bytes: bytes, bbox, *, padding: float = 6.0, upscale: float = 1.0) -> bytes:
    """The same as `crop_region` for a photo: `bbox` in the image's pixels.

    `upscale` enlarges the crop (2.0 for a second, closer look); the pixels
    themselves are not sharpened or altered.
    """
    import io

    from PIL import Image

    x0, y0, x1, y1 = _checked_box(bbox)
    try:
        with Image.open(io.BytesIO(image_bytes)) as img:
            img.load()
            box = (
                max(0, int(x0 - padding)), max(0, int(y0 - padding)),
                min(img.width, int(x1 + padding + 0.999)), min(img.height, int(y1 + padding + 0.999)),
            )
            if box[2] - box[0] < 1 or box[3] - box[1] < 1:
                raise ExtractionError(
                    f"Region {bbox!r} lies outside the image ({img.width} x {img.height} px).")
            part = img.crop(box).convert("RGB")
            if upscale and upscale != 1.0:
                part = part.resize((max(1, int(part.width * upscale)), max(1, int(part.height * upscale))),
                                   Image.LANCZOS)
            out = io.BytesIO()
            part.save(out, "PNG")
            return out.getvalue()
    except ExtractionError:
        raise
    except Exception as exc:
        raise ExtractionError(f"Could not crop the image: {exc}") from exc


def extract_docx(path: Path | str) -> ExtractedDocument:
    """DOCX text.

    Word documents carry no page geometry, so there are no bounding boxes and
    the source viewer cannot highlight a region — paragraphs are recorded in
    reading order instead, and the whole file is treated as one page.
    """
    import docx

    try:
        document = docx.Document(str(path))
    except Exception as exc:
        raise ExtractionError(f"Could not open the Word document: {exc}") from exc

    page = ExtractedPage(page_number=1, has_native_text=True)
    order = 0
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            page.blocks.append(ExtractedBlock(text=text, bbox=None, reading_order=order))
            order += 1

    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                page.blocks.append(
                    ExtractedBlock(text=" | ".join(cells), bbox=None, reading_order=order)
                )
                order += 1

    page.has_native_text = page.char_count >= _min_chars_for_native_text()
    return ExtractedDocument(pages=[page], page_count=1)


def extract(path: Path | str, *, content_type: str = "") -> ExtractedDocument:
    """Dispatch on file type."""
    suffix = Path(path).suffix.lower()
    if suffix == ".pdf" or "pdf" in content_type:
        return extract_pdf(path)
    if suffix in {".docx", ".doc"} or "word" in content_type:
        return extract_docx(path)
    raise ExtractionError(f"Unsupported file type '{suffix}'.")
